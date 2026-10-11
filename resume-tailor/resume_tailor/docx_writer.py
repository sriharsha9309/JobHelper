import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from .engine import TailoredResume
from .master import month_year

DEFAULT_TEMPLATE = {
    "font": "Calibri",
    "body_size": 10.5,
    "name_size": 20,
    "heading_size": 11.5,
    "accent": "1F3864",
    "page": "A4",
    "margin_in": 0.6,
}
_MUTED = RGBColor(0x40, 0x40, 0x40)
_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
    "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _clean(value) -> str:
    return " ".join(str(value or "").split())


def _link_label(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url).rstrip("/")


def _contact_parts(contact: dict) -> list[str]:
    return [_clean(contact.get(k)) for k in ("email", "phone", "location") if _clean(contact.get(k))]


def _contact_line(contact: dict) -> str:
    links = [_link_label(_clean(u)) for u in contact.get("links") or [] if _clean(u)]
    return " | ".join(_contact_parts(contact) + links)


def _dates(item) -> str:
    get = item.get if isinstance(item, dict) else (lambda key: getattr(item, key, ""))
    start, end = _clean(get("start")), _clean(get("end"))
    if not start:
        return month_year(end) if end else ""
    return f"{month_year(start)} - {month_year(end)}" if end else month_year(start)


def file_name(resume: TailoredResume, company: str) -> str:
    parts = [resume.contact.get("name", "Resume"), resume.headline, company]
    stem = re.sub(r"[^A-Za-z0-9]+", "_", " ".join(_clean(p) for p in parts if _clean(p))).strip("_")
    return f"{stem[:120] or 'Resume'}.docx"


class _Builder:
    """Single column, real text only: no tables, text boxes, images, headers or footers (ATS-safe)."""

    def __init__(self, template: dict | None):
        t = {**DEFAULT_TEMPLATE, **(template or {})}
        self.t = t
        self.accent = RGBColor.from_string(str(t["accent"]).lstrip("#").upper())
        self.doc = Document()
        section = self.doc.sections[0]
        if str(t["page"]).lower() == "letter":
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        else:
            section.page_width, section.page_height = Mm(210), Mm(297)
        margin = Inches(float(t["margin_in"]))
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = margin
        self.text_width = section.page_width - section.left_margin - section.right_margin

        normal = self.doc.styles["Normal"]
        normal.font.name = t["font"]
        normal.font.size = Pt(float(t["body_size"]))
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), t["font"])
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(1)
        normal.paragraph_format.line_spacing = 1.05

        bullets = self.doc.styles["List Bullet"]
        bullets.font.name = t["font"]
        bullets.font.size = Pt(float(t["body_size"]))
        bullets.paragraph_format.space_after = Pt(1)

    def para(self, align=None, space_before: float = 0, space_after: float = 1, keep_next: bool = False):
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before, fmt.space_after, fmt.keep_with_next = Pt(space_before), Pt(space_after), keep_next
        if align is not None:
            p.alignment = align
        return p

    def run(self, p, text: str, bold=False, italic=False, size: float | None = None, color: RGBColor | None = None):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return r

    def hyperlink(self, p, url: str, label: str) -> None:
        rel_id = p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), str(self.t["accent"]).lstrip("#"))
        rpr.append(color)
        r.append(rpr)
        text = OxmlElement("w:t")
        text.text = label
        text.set(qn("xml:space"), "preserve")
        r.append(text)
        link.append(r)
        p._p.append(link)

    def right_tab(self, p) -> None:
        p.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)

    def two_sided(self, left: str, right: str, bold=False, italic=False, color=None, space_before: float = 0, keep_next=True):
        p = self.para(space_before=space_before, space_after=0, keep_next=keep_next)
        self.right_tab(p)
        self.run(p, _clean(left), bold=bold, italic=italic, color=color)
        if right:
            self.run(p, "\t" + _clean(right), bold=False, italic=italic, color=color)
        return p

    def heading(self, text: str) -> None:
        p = self.para(space_before=9, space_after=3, keep_next=True)
        self.run(p, text.upper(), bold=True, size=float(self.t["heading_size"]), color=self.accent)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", str(self.t["accent"]).lstrip("#"))):
            bottom.set(qn(key), value)
        border.append(bottom)
        # Word rejects pPr children out of schema order, so place pBdr before its successors.
        p._p.get_or_add_pPr().insert_element_before(border, *_PBDR_SUCCESSORS)

    def bullet(self, text: str) -> None:
        p = self.doc.add_paragraph(_clean(text), style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.16)


def write_docx(r: TailoredResume, path: Path, template: dict | None = None) -> Path:
    b = _Builder(template)
    doc = b.doc
    name = _clean(r.contact.get("name", ""))
    doc.core_properties.author = name
    doc.core_properties.title = f"{name} - {_clean(r.headline)}"

    p = b.para(WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    b.run(p, name, bold=True, size=float(b.t["name_size"]), color=b.accent)
    p = b.para(WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    b.run(p, _clean(r.headline), bold=True, size=float(b.t["heading_size"]))
    p = b.para(WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    b.run(p, " | ".join(_contact_parts(r.contact)), size=float(b.t["body_size"]) - 0.5)
    for url in (u for u in r.contact.get("links") or [] if _clean(u).startswith(("https://", "http://"))):
        b.run(p, " | ", size=float(b.t["body_size"]) - 0.5)
        b.hyperlink(p, _clean(url), _link_label(_clean(url)))

    b.heading("Summary")
    b.run(b.para(space_after=0), " ".join(_clean(s) for s in r.summary))

    b.heading("Skills")
    for group, items in r.skills.items():
        p = b.para(space_after=1)
        b.run(p, f"{_clean(group)}: ", bold=True)
        b.run(p, ", ".join(_clean(i) for i in items))

    b.heading("Work Experience")
    for i, (role, bullets) in enumerate(r.experience):
        b.two_sided(role.company, role.location, bold=True, space_before=0 if i == 0 else 5)
        b.two_sided(role.title, _dates(role), italic=True, color=_MUTED)
        for text in bullets:
            b.bullet(text)

    if r.projects:
        b.heading("Projects")
        for i, (project, bullets) in enumerate(r.projects):
            p = b.para(space_before=0 if i == 0 else 4, space_after=0, keep_next=True)
            b.right_tab(p)
            b.run(p, _clean(project.name), bold=True)
            if project.tech:
                b.run(p, " | " + ", ".join(_clean(t) for t in project.tech), italic=True, color=_MUTED)
            if _clean(project.link).startswith(("https://", "http://")):
                b.run(p, "\t")
                b.hyperlink(p, _clean(project.link), _link_label(_clean(project.link)))
            for text in bullets:
                b.bullet(text)

    if r.education:
        b.heading("Education")
        for e in r.education:
            b.two_sided(e.get("degree", ""), _dates(e), bold=True)
            details = " | ".join(x for x in (_clean(e.get("institution")), _clean(e.get("location"))) if x)
            b.two_sided(details, _clean(e.get("grade")), italic=True, color=_MUTED, keep_next=False)

    if r.certifications:
        b.heading("Certifications")
        for c in r.certifications:
            b.bullet(c)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def render_text(r: TailoredResume) -> str:
    """Plain-text preview - roughly what an ATS parser extracts from the DOCX."""
    out = [_clean(r.contact.get("name", "")), _clean(r.headline), _contact_line(r.contact), "", "SUMMARY", " ".join(r.summary), "", "SKILLS"]
    out += [f"{g}: {', '.join(items)}" for g, items in r.skills.items()]
    out += ["", "WORK EXPERIENCE"]
    for role, bullets in r.experience:
        out += [f"{role.company}    {role.location}", f"{role.title}    {_dates(role)}"]
        out += [f"  - {b}" for b in bullets]
    if r.projects:
        out += ["", "PROJECTS"]
        for p, bullets in r.projects:
            out.append(p.name + (" | " + ", ".join(p.tech) if p.tech else ""))
            out += [f"  - {b}" for b in bullets]
    if r.education:
        out += ["", "EDUCATION"]
        for e in r.education:
            out += [f"{_clean(e.get('degree'))}    {_dates(e)}", " | ".join(x for x in (_clean(e.get("institution")), _clean(e.get("grade"))) if x)]
    if r.certifications:
        out += ["", "CERTIFICATIONS"] + [f"  - {c}" for c in r.certifications]
    return "\n".join(out)
