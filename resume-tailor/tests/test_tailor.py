from collections import Counter

import pytest
from docx import Document

from resume_tailor.config import ROOT, load_config
from resume_tailor.docx_writer import render_text, write_docx
from resume_tailor.engine import apply_ai, check_rewrite, keyword_tailor, tailor
from resume_tailor.keywords import build_vocab, jd_keywords
from resume_tailor.master import load_master, month_year, parse_master
from resume_tailor.server import load_job, request_allowed

JD = (
    "We are hiring a GenAI Engineer to build agentic RAG systems with LangGraph and LangChain on AWS Bedrock. "
    "You will write Python, design prompt engineering strategies, deploy with Docker and Kubernetes, and evaluate LLMs. "
    "Experience with vector databases such as Qdrant is a plus. 1-3 years of experience required."
)

MASTER = {
    "contact": {"name": "Test User", "email": "t@example.com", "phone": "+91 90000 00000", "location": "Pune, India"},
    "headline": "GenAI Engineer",
    "years": "2.5+",
    "summary_points": ["Built production LangGraph agents on AWS.", "Strong at debugging with AppDynamics."],
    "skills": {"GenAI": ["LangGraph", "LangChain", "RAG", "AWS Bedrock"], "Languages": ["Python", "SQL", "Java"]},
    "experience": [
        {
            "id": "r1", "company": "TCS", "title": "Systems Engineer", "location": "Pune", "start": "2024-06", "end": "present",
            "min_bullets": 2, "max_bullets": 3,
            "bullets": [
                {"id": "b1", "text": "Built a LangGraph agent workflow converting 120 legacy files to Java.", "tags": ["GenAI", "agentic"]},
                {"id": "b2", "text": "Implemented RAG over legacy code with LangChain.", "tags": ["RAG"]},
                {"id": "b3", "text": "Monitored apps with AppDynamics on AWS."},
                {"id": "b4", "text": "Improved <metric> by <X>%."},
            ],
        }
    ],
    "projects": [],
    "education": [{"degree": "B.Tech", "institution": "College", "start": "2019-08", "end": "2023-05"}],
    "certifications": ["AWS Certified Data Engineer - Associate"],
}


@pytest.fixture(scope="module")
def cfg():
    return load_config()


@pytest.fixture()
def master():
    return parse_master(MASTER)


def test_month_year():
    assert month_year("2024-06") == "Jun 2024"
    assert month_year("present") == "Present"
    assert month_year("Jun 2024") == "Jun 2024"


def test_example_master_parses_and_reports_placeholders():
    m = load_master(ROOT / "resume" / "master.example.yaml")
    assert m.experience and m.placeholders()


def test_duplicate_bullet_ids_rejected():
    data = {**MASTER, "experience": [{**MASTER["experience"][0], "bullets": [{"id": "x", "text": "a"}, {"id": "x", "text": "b"}]}]}
    with pytest.raises(ValueError):
        parse_master(data)


def test_check_rewrite_guards(cfg, master):
    vocab = build_vocab(cfg["skills"], master)
    src = "Built a LangGraph agent workflow converting 120 legacy files to Java. GenAI agentic"
    assert check_rewrite("Built GenAI LangGraph agents that converted 120 legacy files to Java.", src, vocab) is None
    assert "numbers" in check_rewrite("Built LangGraph agents converting 500 files.", src, vocab)
    assert "skills" in check_rewrite("Built LangGraph agents on Kubernetes converting 120 files.", src, vocab)


def test_keyword_tailor_ranks_and_skips_placeholders(cfg, master):
    vocab = build_vocab(cfg["skills"], master)
    jdk = jd_keywords(JD, vocab)
    resume, selection = keyword_tailor(master, jdk, vocab, "GenAI Engineer", {})
    ids = [b.id for b in selection["r1"]]
    assert "b4" not in ids
    assert 2 <= len(ids) <= 3
    assert ids[0] in ("b1", "b2")
    first_group = next(iter(resume.skills))
    assert first_group == "GenAI"


def test_apply_ai_rejects_invented_facts(cfg, master):
    vocab = build_vocab(cfg["skills"], master)
    jdk = jd_keywords(JD, vocab)
    base, selection = keyword_tailor(master, jdk, vocab, "GenAI Engineer", {})
    data = {
        "summary": ["GenAI Engineer with 2.5+ years building LangGraph agents.", "Expert in Kubernetes with 10 years."],
        "bullets": [
            {"source_id": "b1", "text": "Built agentic GenAI workflows with LangGraph converting 999 legacy files."},
            {"source_id": "nope", "text": "Invented bullet."},
            {"source_id": "b2", "text": "Implemented RAG pipelines over legacy code using LangChain."},
        ],
        "skills_priority": ["Python", "Kubernetes", "LangGraph"],
    }
    warnings = []
    out = apply_ai(master, base, selection, data, vocab, {}, warnings)
    bullets = out.experience[0][1]
    assert bullets[0] == MASTER["experience"][0]["bullets"][0]["text"]
    assert "Implemented RAG pipelines over legacy code using LangChain." in bullets
    assert not any("Kubernetes" in s for s in out.summary)
    assert all("Kubernetes" not in items for items in out.skills.values())
    assert any("unknown bullet" in w for w in warnings)


def test_tailor_without_ai_key(cfg, master, monkeypatch, tmp_path):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    result = tailor(master, JD, "GenAI Engineer", "Acme", cfg, use_exact_title=True, use_ai=True)
    assert result.mode == "keywords"
    assert result.resume.headline == "GenAI Engineer"
    assert result.coverage_after >= result.coverage_before
    assert "kubernetes" in result.not_in_master

    path = write_docx(result.resume, tmp_path / "r.docx")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert not doc.tables
    assert all(not p.text for s in doc.sections for p in s.header.paragraphs)
    for heading in ("SUMMARY", "SKILLS", "WORK EXPERIENCE", "EDUCATION", "CERTIFICATIONS"):
        assert heading in text
    assert "t@example.com" in text and "Jun 2024 - Present" in text
    assert "WORK EXPERIENCE" in render_text(result.resume)


@pytest.mark.parametrize(
    "host, origin, ok",
    [
        ("127.0.0.1:8765", None, True),
        ("localhost:8765", "http://localhost:8765", True),
        ("127.0.0.1:8765", "https://evil.example", False),
        ("evil.example:8765", None, False),
    ],
)
def test_request_allowed(host, origin, ok):
    assert request_allowed(host, origin, 8765) is ok


def test_jd_keywords_counts(cfg, master):
    vocab = build_vocab(cfg["skills"], master)
    jdk = jd_keywords(JD, vocab)
    assert isinstance(jdk, Counter) and jdk["langgraph"] == 1


def test_template_is_ats_safe_and_styled(cfg, monkeypatch, tmp_path):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    data = {**MASTER, "contact": {**MASTER["contact"], "links": ["https://github.com/example"]}}
    result = tailor(parse_master(data), JD, "GenAI Engineer", "Acme", cfg, use_ai=False)
    path = write_docx(result.resume, tmp_path / "styled.docx", cfg.get("template"))
    doc = Document(str(path))
    section = doc.sections[0]
    assert abs(section.page_width.mm - 210) < 1
    assert not doc.tables and not doc.inline_shapes
    assert any("github.com/example" in rel.target_ref for rel in doc.part.rels.values() if rel.is_external)
    company_line = next(p for p in doc.paragraphs if p.text.startswith("TCS"))
    assert "\t" in company_line.text and company_line.paragraph_format.tab_stops[0] is not None
    heading = next(p for p in doc.paragraphs if p.text == "WORK EXPERIENCE")
    assert heading._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr") is not None


def test_real_master_loads_if_present():
    path = ROOT / "resume" / "master.yaml"
    if not path.exists():
        pytest.skip("no personal master.yaml")
    m = load_master(path)
    assert m.experience and m.contact.get("name")


def test_load_job_reads_job_alerts_history(tmp_path):
    jobs = tmp_path / "jobs.json"
    jobs.write_text('{"k1": {"title": "AI Engineer", "company": "Acme", "description": "JD text", "url": "https://x"}}', encoding="utf-8")
    assert load_job("k1", jobs)["description"] == "JD text"
    assert load_job("missing", jobs) == {"title": "", "company": "", "url": "", "description": ""}
    assert load_job("k1", tmp_path / "nope.json") == {}
